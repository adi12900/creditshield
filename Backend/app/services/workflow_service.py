from __future__ import annotations

import io
import os
import re
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any

import boto3
import httpx
import pdfplumber
from botocore.exceptions import BotoCoreError, ClientError
from openpyxl import load_workbook

from app.core.database import SessionLocal
from app.models.document import Document
from app.models.loan_application import LoanApplication
from app.models.user import User
from app.services.s3 import download_bytes_from_s3
from app.schemas.workflow import AuditLogItem, RegulatoryReport


class WorkflowServiceError(Exception):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


class WorkflowService:
    def __init__(self) -> None:
        self._applications: dict[str, dict[str, Any]] = {
            "ARN202600001": {
                "arn": "ARN202600001",
                "borrower_name": "Rajesh Kumar",
                "loan_amount": 500000,
                "stage": "Lead",
                "risk_grade": "A+",
                "credit_score": 730,
                "kyc_status": "Pending",
                "employment_type": "Salaried",
                "purpose": "Home Renovation",
            },
            "ARN202600004": {
                "arn": "ARN202600004",
                "borrower_name": "Vikram Singh",
                "loan_amount": 850000,
                "stage": "Submitted",
                "risk_grade": "A+",
                "credit_score": 720,
                "kyc_status": "Verified",
                "employment_type": "Salaried",
                "purpose": "Business Expansion",
            },
            "ARN202600005": {
                "arn": "ARN202600005",
                "borrower_name": "Neha Gupta",
                "loan_amount": 650000,
                "stage": "Documents Pending",
                "risk_grade": "B",
                "credit_score": 680,
                "kyc_status": "Pending",
                "employment_type": "Salaried",
                "purpose": "Medical Expenses",
            },
            "ARN202600009": {
                "arn": "ARN202600009",
                "borrower_name": "Rahul Verma",
                "loan_amount": 800000,
                "stage": "Underwriting",
                "risk_grade": "A+",
                "credit_score": 735,
                "kyc_status": "Verified",
                "employment_type": "Self Employed",
                "purpose": "Inventory Funding",
            },
            "ARN202600011": {
                "arn": "ARN202600011",
                "borrower_name": "Suresh Rao",
                "loan_amount": 1200000,
                "stage": "Offer Sent",
                "risk_grade": "A",
                "credit_score": 715,
                "kyc_status": "Verified",
                "employment_type": "Salaried",
                "purpose": "Debt Consolidation",
            },
            "ARN202600015": {
                "arn": "ARN202600015",
                "borrower_name": "Anil Kumar",
                "loan_amount": 1100000,
                "stage": "Rejected",
                "risk_grade": "C",
                "credit_score": 640,
                "kyc_status": "Verified",
                "employment_type": "Self Employed",
                "purpose": "Working Capital",
            },
        }

        self._documents: dict[str, list[dict[str, Any]]] = {
            arn: [
                {"id": "901", "type": "Aadhaar Card", "status": "Verified", "confidence": 98},
                {"id": "902", "type": "PAN Card", "status": "Verified", "confidence": 96},
                {"id": "903", "type": "Bank Statement", "status": "Verified", "confidence": 94},
            ]
            for arn in self._applications
        }

        self._communications: dict[str, list[dict[str, Any]]] = {
            arn: [
                {
                    "id": "c1",
                    "channel": "email",
                    "subject": "Application Received",
                    "message": f"Thank you for applying. ARN: {arn}",
                    "sent_at": datetime.now(tz=timezone.utc).isoformat(),
                }
            ]
            for arn in self._applications
        }

        self._credit_memos: dict[str, dict[str, Any]] = {}
        self._policy_overrides: dict[str, list[dict[str, Any]]] = {}
        self._document_ocr_cache: dict[str, dict[str, Any]] = {}

        self._reports: list[RegulatoryReport] = [
            RegulatoryReport(
                id="REP-001",
                name="RBI NBFC Returns - Q1 2026",
                report_type="Quarterly",
                due_date="2026-04-15",
                status="In Progress",
                completeness=85,
            ),
            RegulatoryReport(
                id="REP-002",
                name="AML/KYC Compliance Summary",
                report_type="Monthly",
                due_date="2026-04-05",
                status="Submitted",
                completeness=100,
            ),
        ]

        self._audit_logs: list[AuditLogItem] = [
            AuditLogItem(
                timestamp=datetime.now(tz=timezone.utc),
                user="System",
                action="Workflow service initialized",
                resource="workflow",
                details="Role APIs bootstrapped",
                risk="Low",
            )
        ]

        self._allowed_stages = {
            "Lead",
            "Submitted",
            "Documents Pending",
            "KYC",
            "Underwriting",
            "Offer Sent",
            "Disbursed",
            "Rejected",
        }
        self._allowed_risk_grades = {"A+", "A", "B", "C"}
        self._allowed_kyc_statuses = {"Verified", "Pending"}
        self._allowed_employment_types = {"Salaried", "Self Employed"}

    def _to_workflow_application(self, row: LoanApplication) -> dict[str, Any]:
        stage = row.stage if row.stage in self._allowed_stages else "Submitted"
        risk_grade = row.risk_grade if row.risk_grade in self._allowed_risk_grades else "B"
        kyc_status = row.kyc_status if row.kyc_status in self._allowed_kyc_statuses else "Pending"
        employment_type = row.employment_type if row.employment_type in self._allowed_employment_types else "Salaried"
        loan_amount = float(row.loan_amount if isinstance(row.loan_amount, Decimal) else row.loan_amount or 0)

        return {
            "arn": row.arn,
            "borrower_name": row.borrower_name,
            "borrower_email": row.borrower_email,
            "borrower_phone": row.borrower_phone,
            "loan_amount": loan_amount,
            "loan_type": row.loan_type or "digital_personal_loan",
            "stage": stage,
            "risk_grade": risk_grade,
            "credit_score": int(row.credit_score or 700),
            "kyc_status": kyc_status,
            "employment_type": employment_type,
            "purpose": row.purpose or "General Purpose",
            "created_at": row.created_at,
            "updated_at": row.updated_at,
        }

    def _update_application_fields(self, arn: str, **fields: Any) -> bool:
        try:
            with SessionLocal() as db:
                row = db.query(LoanApplication).filter(LoanApplication.arn == arn).first()
                if not row:
                    return False
                for field, value in fields.items():
                    if hasattr(row, field):
                        setattr(row, field, value)
                db.add(row)
                db.commit()
                return True
        except Exception:
            return False

    def _get_application(self, arn: str) -> dict[str, Any]:
        try:
            with SessionLocal() as db:
                row = db.query(LoanApplication).filter(LoanApplication.arn == arn).first()
                if row:
                    return self._to_workflow_application(row)
        except Exception:
            # Fallback to in-memory dataset when DB is unavailable.
            pass

        app = self._applications.get(arn)
        if not app:
            raise WorkflowServiceError(f"Application not found for ARN {arn}", 404)
        return app

    def list_applications(self, stage: str | None = None) -> list[dict[str, Any]]:
        try:
            with SessionLocal() as db:
                query = db.query(LoanApplication)
                if stage:
                    query = query.filter(LoanApplication.stage == stage)
                rows = query.order_by(LoanApplication.id.asc()).all()
                if rows:
                    return [self._to_workflow_application(row) for row in rows]
        except Exception:
            # Fallback to in-memory dataset when DB is unavailable.
            pass

        applications = list(self._applications.values())
        if stage:
            applications = [app for app in applications if app["stage"] == stage]
        return applications

    def get_application(self, arn: str) -> dict[str, Any]:
        return self._get_application(arn)

    def role_dashboard(self, role: str) -> dict[str, Any]:
        apps = self.list_applications()
        if role == "loan_officer":
            return {
                "role": role,
                "stats": [
                    {"key": "active_applications", "value": len([a for a in apps if a["stage"] not in {"Disbursed", "Rejected"}])},
                    {"key": "sla_breaches", "value": len([a for a in apps if a["stage"] == "Documents Pending"])},
                ],
            }
        if role == "credit_analyst":
            return {
                "role": role,
                "stats": [
                    {"key": "in_progress", "value": len([a for a in apps if a["stage"] in {"Submitted", "Documents Pending"}])},
                    {"key": "avg_credit_score", "value": int(sum(a["credit_score"] for a in apps) / len(apps))},
                ],
            }
        if role == "underwriter":
            return {
                "role": role,
                "stats": [
                    {"key": "underwriting_queue", "value": len([a for a in apps if a["stage"] == "Underwriting"])},
                    {"key": "approval_rate", "value": "78%"},
                ],
            }
        return {
            "role": role,
            "stats": [
                {"key": "kyc_pending", "value": len([a for a in apps if a["kyc_status"] == "Pending"])},
                {"key": "reports_due", "value": len([r for r in self._reports if r.status != "Submitted"])},
            ],
        }

    def get_documents(self, arn: str) -> list[dict[str, Any]]:
        self._get_application(arn)

        try:
            with SessionLocal() as db:
                rows = (
                    db.query(Document)
                    .join(LoanApplication, LoanApplication.id == Document.application_id)
                    .filter(LoanApplication.arn == arn)
                    .order_by(Document.uploaded_at.desc(), Document.id.desc())
                    .all()
                )

                return [
                    {
                        "id": str(row.id),
                        "type": row.doc_type,
                        "status": row.status,
                        "confidence": int(row.confidence) if row.confidence is not None else None,
                        "storage_url": row.storage_url,
                        "uploaded_by_user_id": row.uploaded_by_user_id,
                        "uploaded_at": row.uploaded_at,
                    }
                    for row in rows
                ]
        except Exception:
            # Fallback to in-memory dataset when DB is unavailable.
            pass

        return self._documents.get(arn, [])

    def _normalize_doc_type(self, value: str) -> str:
        return value.lower().replace(" ", "_").replace("-", "_")

    def _ocr_cache_key(self, arn: str, document_id: str) -> str:
        return f"{arn}:{document_id}"

    def _download_document_bytes(self, storage_url: str) -> bytes:
        try:
            response = httpx.get(storage_url, timeout=20.0, follow_redirects=True)
            response.raise_for_status()
            return response.content
        except Exception:
            if ".s3." in storage_url:
                try:
                    return download_bytes_from_s3(storage_url)
                except Exception as exc:
                    raise WorkflowServiceError(f"Unable to download document from storage: {exc}", 502) from exc
            raise WorkflowServiceError("Unable to download document for OCR extraction", 502)

    def _extract_text_with_textract(self, document_bytes: bytes) -> str:
        try:
            client = boto3.client(
                "textract",
                aws_access_key_id=os.getenv("AWS_ACCESS_KEY"),
                aws_secret_access_key=os.getenv("AWS_SECRET_KEY"),
                region_name=os.getenv("AWS_REGION") or os.getenv("AWS_REGIONS3") or "ap-south-1",
            )
            response = client.detect_document_text(Document={"Bytes": document_bytes})
            lines = [
                block.get("Text", "")
                for block in response.get("Blocks", [])
                if block.get("BlockType") == "LINE" and block.get("Text")
            ]
            return "\n".join(lines).strip()
        except ClientError as exc:
            code = exc.response.get("Error", {}).get("Code", "")
            if code == "SubscriptionRequiredException":
                # Graceful fallback when Textract is not enabled for this AWS account.
                return ""
            raise WorkflowServiceError(f"Textract OCR failed: {exc}", 502) from exc
        except BotoCoreError as exc:
            raise WorkflowServiceError(f"Textract OCR failed: {exc}", 502) from exc

    def _build_fallback_ocr_data(self, app: dict[str, Any], doc_type: str) -> dict[str, dict[str, Any]]:
        borrower_name = str(app.get("borrower_name") or "Not found")
        return {
            "Document Type": {"value": doc_type, "confidence": 70},
            "Borrower": {"value": borrower_name, "confidence": 72},
            "Extraction Note": {
                "value": "Primary OCR engine unavailable for this file/account. Parsed basic metadata fallback.",
                "confidence": 60,
            },
        }

    def _extract_text_from_pdf(self, pdf_bytes: bytes) -> str:
        text_chunks: list[str] = []
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    text_chunks.append(text)

        return "\n".join(text_chunks).strip()

    def _extract_text_from_xlsx(self, xlsx_bytes: bytes) -> str:
        workbook = load_workbook(filename=io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
        values: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    values.append(" ".join(cells))
        return "\n".join(values).strip()

    def _extract_raw_text(self, storage_url: str) -> str:
        lower_url = storage_url.lower()
        file_bytes = self._download_document_bytes(storage_url)

        if lower_url.endswith(".pdf"):
            pdf_text = self._extract_text_from_pdf(file_bytes)
            if pdf_text:
                return pdf_text
            return self._extract_text_with_textract(file_bytes)

        if lower_url.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff")):
            return self._extract_text_with_textract(file_bytes)

        if lower_url.endswith(".csv"):
            return file_bytes.decode("utf-8", errors="ignore")

        if lower_url.endswith(".xlsx"):
            return self._extract_text_from_xlsx(file_bytes)

        return file_bytes.decode("utf-8", errors="ignore")

    def _extract_by_pattern(self, text: str, pattern: str) -> str | None:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if not match:
            return None
        value = match.group(1) if match.lastindex else match.group(0)
        cleaned = re.sub(r"\s+", " ", value).strip()
        return cleaned or None

    def _field(self, value: str | None, fallback: str, confidence_if_found: int, confidence_if_missing: int = 55) -> dict[str, Any]:
        if value:
            return {"value": value, "confidence": confidence_if_found}
        return {"value": fallback, "confidence": confidence_if_missing}

    def _build_ocr_data_from_text(self, app: dict[str, Any], doc_type: str, raw_text: str) -> dict[str, dict[str, Any]]:
        normalized = self._normalize_doc_type(doc_type)
        compact_text = re.sub(r"\s+", " ", raw_text).strip()
        borrower_name = str(app.get("borrower_name") or "").upper()

        aadhaar = self._extract_by_pattern(compact_text, r"\b(\d{4}\s?\d{4}\s?\d{4})\b")
        pan = self._extract_by_pattern(compact_text, r"\b([A-Z]{5}\d{4}[A-Z])\b")
        dob = self._extract_by_pattern(compact_text, r"(?:DOB|Date of Birth)\s*[:\-]?\s*([0-9]{1,2}[\-/][0-9]{1,2}[\-/][0-9]{2,4})")
        account_no = self._extract_by_pattern(compact_text, r"(?:A\/C|Account)\s*(?:No|Number)?\s*[:\-]?\s*([0-9Xx]{6,20})")
        ifsc = self._extract_by_pattern(compact_text, r"\b([A-Z]{4}0[A-Z0-9]{6})\b")
        income = self._extract_by_pattern(compact_text, r"(?:Net\s*Salary|Income|Net\s*Pay)\s*[:\-]?\s*([\u20b9RsINR\.,\s0-9]+)")

        if "aadhaar" in normalized:
            return {
                "Aadhaar Number": self._field(aadhaar, "Not found", 96),
                "Name": self._field(borrower_name or None, "Not found", 90),
                "DOB": self._field(dob, "Not found", 88),
            }

        if "pan" in normalized:
            return {
                "PAN Number": self._field(pan, "Not found", 97),
                "Name": self._field(borrower_name or None, "Not found", 90),
                "DOB": self._field(dob, "Not found", 85),
            }

        if "bank_statement" in normalized:
            return {
                "Account Number": self._field(account_no, "Not found", 90),
                "IFSC": self._field(ifsc, "Not found", 86),
                "Borrower": self._field(borrower_name or None, "Not found", 84),
            }

        if "salary" in normalized or "income_proof" in normalized:
            return {
                "Borrower": self._field(borrower_name or None, "Not found", 90),
                "Net Income": self._field(income, "Not found", 88),
            }

        snippet = compact_text[:180] + ("..." if len(compact_text) > 180 else "")
        return {
            "Document Type": self._field(doc_type, "Not found", 90),
            "Extracted Text Snippet": self._field(snippet or None, "No OCR text extracted", 80),
        }

    def extract_document_ocr(self, arn: str, document_id: str) -> dict[str, Any]:
        app = self._get_application(arn)
        docs = self.get_documents(arn)
        target = next((doc for doc in docs if str(doc.get("id")) == str(document_id)), None)
        if not target:
            raise WorkflowServiceError(f"Document {document_id} not found", 404)

        storage_url = str(target.get("storage_url") or "").strip()
        if not storage_url:
            raise WorkflowServiceError("Document has no storage URL. Upload file before OCR extraction.", 400)

        raw_text = self._extract_raw_text(storage_url)
        doc_type = str(target.get("type") or "Document")
        ocr_data = (
            self._build_ocr_data_from_text(app, doc_type, raw_text)
            if raw_text.strip()
            else self._build_fallback_ocr_data(app, doc_type)
        )
        extracted_at = datetime.now(tz=timezone.utc)

        confidences = [int(field.get("confidence", 0)) for field in ocr_data.values()]
        derived_confidence = round(sum(confidences) / len(confidences)) if confidences else None

        try:
            with SessionLocal() as db:
                doc_id_int = int(document_id)
                row = (
                    db.query(Document)
                    .join(LoanApplication, LoanApplication.id == Document.application_id)
                    .filter(LoanApplication.arn == arn, Document.id == doc_id_int)
                    .first()
                )
                if row is not None and derived_confidence is not None:
                    row.confidence = derived_confidence
                    db.add(row)
                    db.commit()
        except Exception:
            # Confidence persistence failure should not block OCR response.
            pass

        response = {
            "document_id": str(target.get("id")),
            "arn": arn,
            "doc_type": doc_type,
            "ocr_data": ocr_data,
            "extracted_at": extracted_at,
        }
        self._document_ocr_cache[self._ocr_cache_key(arn, str(document_id))] = response
        return response

    def get_document_ocr(self, arn: str, document_id: str) -> dict[str, Any]:
        cache_key = self._ocr_cache_key(arn, str(document_id))
        cached = self._document_ocr_cache.get(cache_key)
        if cached:
            return cached

        raise WorkflowServiceError("OCR data not available. Run extraction first.", 404)

    def review_document(self, arn: str, document_id: str, decision: str, reason: str | None) -> dict[str, Any]:
        next_status = "Verified" if decision == "approve" else "Flagged"

        try:
            with SessionLocal() as db:
                try:
                    doc_id_int = int(document_id)
                except ValueError as exc:
                    raise WorkflowServiceError(f"Document {document_id} not found", 404) from exc

                target = (
                    db.query(Document)
                    .join(LoanApplication, LoanApplication.id == Document.application_id)
                    .filter(
                        LoanApplication.arn == arn,
                        Document.id == doc_id_int,
                    )
                    .first()
                )
                if target:
                    target.status = next_status
                    db.add(target)
                    db.commit()
                    db.refresh(target)

                    self.add_audit_log(
                        user="Loan Officer",
                        action=f"Document {decision.title()}",
                        resource=f"{arn}:{document_id}",
                        details=reason or "No reason provided",
                        risk="Low" if decision == "approve" else "Medium",
                    )
                    return {
                        "id": str(target.id),
                        "type": target.doc_type,
                        "status": target.status,
                        "confidence": int(target.confidence) if target.confidence is not None else None,
                        "storage_url": target.storage_url,
                        "uploaded_by_user_id": target.uploaded_by_user_id,
                        "uploaded_at": target.uploaded_at,
                    }
        except WorkflowServiceError:
            raise
        except Exception:
            # Fallback to in-memory dataset when DB is unavailable.
            pass

        docs = self.get_documents(arn)
        target = next((doc for doc in docs if doc["id"] == document_id), None)
        if not target:
            raise WorkflowServiceError(f"Document {document_id} not found", 404)

        target["status"] = next_status
        self.add_audit_log(
            user="Loan Officer",
            action=f"Document {decision.title()}",
            resource=f"{arn}:{document_id}",
            details=reason or "No reason provided",
            risk="Low" if decision == "approve" else "Medium",
        )
        return target

    def get_communications(self, arn: str) -> list[dict[str, Any]]:
        self._get_application(arn)
        return self._communications.get(arn, [])

    def send_communication(self, arn: str, channel: str, subject: str, message: str) -> dict[str, Any]:
        self._get_application(arn)
        history = self._communications.setdefault(arn, [])
        item = {
            "id": f"c{len(history) + 1}",
            "channel": channel,
            "subject": subject,
            "message": message,
            "sent_at": datetime.now(tz=timezone.utc).isoformat(),
        }
        history.insert(0, item)
        self.add_audit_log(
            user="Loan Officer",
            action="Communication Sent",
            resource=arn,
            details=f"{channel.upper()} - {subject}",
            risk="Low",
        )
        return item

    def move_stage(self, arn: str, new_stage: str, action: str, actor: str) -> dict[str, Any]:
        app = self._get_application(arn)
        previous = app["stage"]
        updated_db = self._update_application_fields(arn, stage=new_stage)
        if not updated_db:
            app["stage"] = new_stage
        self.add_audit_log(
            user=actor,
            action=action,
            resource=arn,
            details=f"Stage moved from {previous} to {new_stage}",
            risk="Low",
        )
        return self._get_application(arn)

    def submit_to_credit_analyst(self, arn: str, *, file_complete: bool, actor: str) -> dict[str, Any]:
        app = self._get_application(arn)
        stage = app["stage"]

        # Idempotent path: if already beyond intake submission, do not fail the action.
        if stage in {"Documents Pending", "Underwriting", "Offer Sent", "Disbursed"}:
            self.add_audit_log(
                user=actor,
                action="Submit to Credit Analyst (No-op)",
                resource=arn,
                details=f"Application already progressed to stage {stage}",
                risk="Low",
            )
            return app

        if stage != "Submitted":
            raise WorkflowServiceError("Application can be submitted to credit analyst only from Submitted stage", 400)

        if not file_complete:
            raise WorkflowServiceError("Please confirm that applicant file is complete before submitting", 400)

        docs = self.get_documents(arn)
        if not docs:
            raise WorkflowServiceError("No documents found. Upload and verify documents before submission", 400)

        not_verified = [doc for doc in docs if doc.get("status") != "Verified"]
        if not_verified:
            raise WorkflowServiceError("All documents must be verified before submission to credit analyst", 400)

        rbi_status = self.get_rbi_compliance(arn)
        if int(rbi_status.get("blocking_issues", 0)) > 0:
            raise WorkflowServiceError("RBI mandatory checks are incomplete. Resolve all blocking issues first", 400)

        updated = self.move_stage(arn, "Documents Pending", "Submitted to Credit Analyst", actor)
        self.add_audit_log(
            user=actor,
            action="Applicant File Marked Complete",
            resource=arn,
            details="Loan officer confirmed file complete and routed to credit analyst queue",
            risk="Low",
        )
        return updated

    def recalculate_ratios(
        self,
        monthly_income: float,
        existing_obligations: float,
        proposed_emi: float,
        loan_amount: float,
        asset_value: float,
    ) -> dict[str, Any]:
        dti = ((existing_obligations + proposed_emi) / monthly_income) * 100
        foir = (existing_obligations / monthly_income) * 100
        ltv = (loan_amount / asset_value) * 100
        return {
            "dti": round(dti, 2),
            "foir": round(foir, 2),
            "ltv": round(ltv, 2),
            "policy_pass": dti <= 45 and foir <= 40 and ltv <= 80,
        }

    def save_credit_memo(self, arn: str, payload: dict[str, Any], submitted: bool) -> dict[str, Any]:
        self._get_application(arn)
        payload = dict(payload)
        payload["submitted"] = submitted
        payload["updated_at"] = datetime.now(tz=timezone.utc).isoformat()
        self._credit_memos[arn] = payload
        if submitted:
            self.move_stage(arn, "Underwriting", "Credit Memo Submitted", "Credit Analyst")
        return payload

    def generate_offer(self, arn: str, loan_amount: float, tenure_months: int, interest_rate: float) -> dict[str, Any]:
        self._get_application(arn)
        monthly_rate = interest_rate / 12 / 100
        emi = (loan_amount * monthly_rate * (1 + monthly_rate) ** tenure_months) / (((1 + monthly_rate) ** tenure_months) - 1)
        total_payable = emi * tenure_months
        total_interest = total_payable - loan_amount
        if not self._update_application_fields(arn, stage="Offer Sent"):
            app = self._get_application(arn)
            app["stage"] = "Offer Sent"
        self.add_audit_log(
            user="Underwriter",
            action="Loan Offer Generated",
            resource=arn,
            details=f"Amount {loan_amount}, tenure {tenure_months}, rate {interest_rate}",
            risk="Low",
        )
        return {
            "arn": arn,
            "emi": round(emi, 2),
            "total_interest": round(total_interest, 2),
            "total_payable": round(total_payable, 2),
        }

    def submit_policy_override(self, arn: str, payload: dict[str, Any]) -> dict[str, Any]:
        self._get_application(arn)
        overrides = self._policy_overrides.setdefault(arn, [])
        item = {
            "id": f"ovr-{len(overrides) + 1}",
            **payload,
            "submitted_at": datetime.now(tz=timezone.utc).isoformat(),
        }
        overrides.append(item)
        self.add_audit_log(
            user="Underwriter",
            action="Policy Override Submitted",
            resource=arn,
            details=payload["override_category"],
            risk="Medium",
        )
        return item

    def get_kyc_aml(self, arn: str) -> dict[str, Any]:
        app = self._get_application(arn)
        return {
            "arn": arn,
            "kyc_status": app["kyc_status"],
            "aml_status": "Clear" if app["risk_grade"] in {"A+", "A", "B"} else "Pending Review",
            "fraud_score": 22 if app["risk_grade"] in {"A+", "A"} else 45,
            "can_clear_hold": app["kyc_status"] == "Verified",
        }

    def clear_compliance_hold(self, arn: str, reason: str) -> dict[str, Any]:
        self._get_application(arn)
        if not self._update_application_fields(arn, kyc_status="Verified"):
            app = self._get_application(arn)
            app["kyc_status"] = "Verified"
        self.add_audit_log(
            user="Compliance Officer",
            action="Compliance Hold Cleared",
            resource=arn,
            details=reason,
            risk="Medium",
        )
        return {"arn": arn, "status": "hold_cleared"}

    def get_fraud_signals(self, arn: str) -> dict[str, Any]:
        app = self._get_application(arn)
        score = 18 if app["risk_grade"] in {"A+", "A"} else 39
        return {
            "arn": arn,
            "aggregate_score": score,
            "risk_band": "Low" if score <= 30 else "Medium",
            "signals": [
                {"label": "Identity Fraud", "score": max(score - 8, 5), "severity": "Low"},
                {"label": "Document Fraud", "score": score, "severity": "Medium" if score > 30 else "Low"},
            ],
        }

    def mark_false_positive(self, arn: str, reason: str) -> dict[str, str]:
        self._get_application(arn)
        self.add_audit_log(
            user="Compliance Officer",
            action="Fraud Signal Marked False Positive",
            resource=arn,
            details=reason,
            risk="Low",
        )
        return {"arn": arn, "status": "false_positive_marked"}

    def list_audit_logs(
        self,
        action: str | None = None,
        user: str | None = None,
        resource: str | None = None,
        risk: str | None = None,
    ) -> list[AuditLogItem]:
        logs = self._audit_logs
        if action:
            logs = [item for item in logs if action.lower() in item.action.lower()]
        if user:
            logs = [item for item in logs if user.lower() in item.user.lower()]
        if resource:
            logs = [item for item in logs if resource.lower() in item.resource.lower()]
        if risk:
            logs = [item for item in logs if item.risk.lower() == risk.lower()]
        return logs

    def add_audit_log(self, user: str, action: str, resource: str, details: str, risk: str) -> None:
        item = AuditLogItem(
            timestamp=datetime.now(tz=timezone.utc),
            user=user,
            action=action,
            resource=resource,
            details=details,
            risk=risk,
        )
        self._audit_logs.insert(0, item)

    def list_reports(self) -> list[RegulatoryReport]:
        return self._reports

    def generate_report(self, name: str, report_type: str, reporting_period: str) -> RegulatoryReport:
        report = RegulatoryReport(
            id=f"REP-{len(self._reports) + 1:03d}",
            name=f"{name} ({reporting_period})",
            report_type=report_type,
            due_date=datetime.now(tz=timezone.utc).date().isoformat(),
            status="In Progress",
            completeness=5,
        )
        self._reports.insert(0, report)
        self.add_audit_log(
            user="Compliance Officer",
            action="Regulatory Report Generated",
            resource=report.id,
            details=report.name,
            risk="Low",
        )
        return report

    def get_rbi_compliance(self, arn: str) -> dict[str, Any]:
        app = self._get_application(arn)
        missing = 1 if app["kyc_status"] == "Pending" else 0
        return {
            "arn": arn,
            "score": 93 if missing == 0 else 81,
            "blocking_issues": missing,
            "items": [
                {
                    "id": "dl-001",
                    "clause": "3.2",
                    "requirement": "KFS disclosure before execution",
                    "status": "Compliant",
                    "value": "Provided",
                },
                {
                    "id": "dl-004",
                    "clause": "5.1",
                    "requirement": "KYC/AML must be verified",
                    "status": "Missing" if missing else "Compliant",
                    "value": app["kyc_status"],
                },
            ],
        }

    def get_rbi_audit_export(self, arn: str) -> dict[str, Any]:
        data = self.get_rbi_compliance(arn)
        return {
            "arn": arn,
            "exported_at": datetime.now(tz=timezone.utc).isoformat(),
            "rows": data["items"],
        }

    def get_lead_prequalification_checklist(self, arn: str) -> list[dict[str, Any]]:
        app = self._get_application(arn)
        return [
            {"id": "lead-001", "item": "Income band captured", "done": bool(app.get("purpose"))},
            {"id": "lead-002", "item": "Employment type captured", "done": bool(app.get("employment_type"))},
            {"id": "lead-003", "item": "Location eligibility validated", "done": True},
            {"id": "lead-004", "item": "Product fit validated", "done": app.get("risk_grade") in {"A+", "A", "B"}},
        ]

    def get_esign_checklist(self, arn: str) -> list[dict[str, Any]]:
        app = self._get_application(arn)
        stage = app.get("stage")
        return [
            {"id": "esign-001", "item": "Offer generated and accepted", "done": stage in {"Offer Sent", "Disbursed"}},
            {"id": "esign-002", "item": "Agreement template populated", "done": stage in {"Offer Sent", "Disbursed"}},
            {"id": "esign-003", "item": "Borrower OTP authentication", "done": stage != "Lead"},
            {"id": "esign-004", "item": "Digital signature captured", "done": stage == "Disbursed"},
            {"id": "esign-005", "item": "Executed copy delivered", "done": stage == "Disbursed"},
        ]

    def get_communication_templates(self) -> list[dict[str, Any]]:
        return [
            {
                "id": "tmpl-001",
                "name": "Document Request",
                "category": "Request",
                "channel": "email",
                "subject": "Additional documents required for your loan application",
                "body": "Dear {borrower_name}, please upload pending documents for ARN {arn} to continue processing.",
            },
            {
                "id": "tmpl-002",
                "name": "Status Update",
                "category": "Update",
                "channel": "email",
                "subject": "Status update for ARN {arn}",
                "body": "Dear {borrower_name}, your application is currently in {stage} stage.",
            },
            {
                "id": "tmpl-003",
                "name": "Offer Notification",
                "category": "Offer",
                "channel": "email",
                "subject": "Loan offer ready for your review",
                "body": "Dear {borrower_name}, your offer has been generated for ARN {arn}. Please review and accept.",
            },
            {
                "id": "tmpl-004",
                "name": "Reminder SMS",
                "category": "Reminder",
                "channel": "sms",
                "subject": "Loan application reminder",
                "body": "Reminder: action required on your application ARN {arn}.",
            },
        ]

    def get_loan_officer_application_summary(self, arn: str) -> dict[str, Any]:
        app = self._get_application(arn)
        docs = self.get_documents(arn)
        comms = self.get_communications(arn)

        email_count = len([item for item in comms if item.get("channel") == "email"])
        sms_count = len([item for item in comms if item.get("channel") == "sms"])
        call_count = len([item for item in comms if item.get("channel") == "call"])

        stage = str(app.get("stage", "Submitted"))
        timeline = [
            {"status": "Submitted", "date": datetime.now(tz=timezone.utc).date().isoformat(), "active": True},
            {"status": "Documents Verified", "date": datetime.now(tz=timezone.utc).date().isoformat(), "active": len(docs) > 0},
            {"status": "KYC Cleared", "date": datetime.now(tz=timezone.utc).date().isoformat(), "active": app.get("kyc_status") == "Verified"},
            {"status": "Credit Analysis", "date": datetime.now(tz=timezone.utc).date().isoformat(), "active": stage in {"Underwriting", "Offer Sent", "Disbursed"}},
            {"status": "Underwriting", "date": datetime.now(tz=timezone.utc).date().isoformat(), "active": stage in {"Underwriting", "Offer Sent", "Disbursed"}},
        ]

        processing_time_days = 0.0
        updated_at = app.get("updated_at")
        if isinstance(updated_at, datetime):
            processing_time_days = round(max((datetime.now(tz=timezone.utc) - updated_at).total_seconds(), 0) / 86400, 1)

        risk_score = int(app.get("credit_score", 700))
        risk_confidence = 92 if app.get("risk_grade") in {"A+", "A"} else 84

        return {
            "arn": arn,
            "application_status": "Completed" if stage == "Disbursed" else "Rejected" if stage == "Rejected" else "In Progress",
            "active_stage": stage,
            "processing_time_days": processing_time_days,
            "documents_verified": len([doc for doc in docs if doc.get("status") == "Verified"]),
            "documents_total": len(docs),
            "communications_total": len(comms),
            "email_count": email_count,
            "sms_count": sms_count,
            "call_count": call_count,
            "risk_score": risk_score,
            "risk_confidence_percent": risk_confidence,
            "timeline": timeline,
        }

    def _list_users(self) -> list[User]:
        try:
            with SessionLocal() as db:
                return db.query(User).order_by(User.id.desc()).all()
        except Exception:
            return []

    def system_admin_dashboard(self) -> dict[str, Any]:
        users = self._list_users()
        apps = self.list_applications()
        active_users = len([u for u in users if u.is_active])
        integrations = [
            {"name": "CIBIL Bureau", "status": "Healthy", "latency_ms": 120},
            {"name": "Experian Bureau", "status": "Healthy", "latency_ms": 150},
            {"name": "eSign Provider", "status": "Healthy", "latency_ms": 200},
            {"name": "Payment Rails (NEFT/RTGS)", "status": "Healthy", "latency_ms": 180},
            {"name": "KYC Provider", "status": "Degraded", "latency_ms": 450},
            {"name": "Core Banking API", "status": "Healthy", "latency_ms": 90},
        ]

        role_counts: dict[str, tuple[int, int]] = {
            "Loan Officers": (0, 0),
            "Credit Analysts": (0, 0),
            "Underwriters": (0, 0),
            "Compliance Officers": (0, 0),
        }

        role_mapping = {
            "loan_officer": "Loan Officers",
            "credit_analyst": "Credit Analysts",
            "underwriter": "Underwriters",
            "compliance_officer": "Compliance Officers",
        }

        for user in users:
            role_label = role_mapping.get(str(user.role))
            if not role_label:
                continue
            current_active, current_total = role_counts[role_label]
            role_counts[role_label] = (
                current_active + (1 if user.is_active else 0),
                current_total + 1,
            )

        role_activity = [
            {"role": role, "active_users": active, "total_users": total}
            for role, (active, total) in role_counts.items()
        ]

        recent_events = [
            {
                "timestamp": item.timestamp,
                "event": f"{item.action}: {item.details}",
                "user": item.user,
            }
            for item in self._audit_logs[:5]
        ]

        return {
            "metrics": [
                {"key": "total_users", "label": "Total Users", "value": len(users)},
                {"key": "active_sessions", "label": "Active Sessions", "value": active_users},
                {"key": "workflows", "label": "Workflows", "value": 1},
                {
                    "key": "integrations",
                    "label": "Integrations",
                    "value": len(integrations),
                    "subtitle": "All healthy" if all(i["status"] == "Healthy" for i in integrations) else "Action required",
                },
            ],
            "integrations": integrations,
            "role_activity": role_activity,
            "recent_events": recent_events,
            "active_applications": len([a for a in apps if a["stage"] not in {"Disbursed", "Rejected"}]),
        }

    def workflow_designer_data(self) -> dict[str, Any]:
        stages = [
            {
                "id": 1,
                "name": "Application Submission",
                "assigned_role": "Loan Officer",
                "avg_duration_minutes": 5,
                "status": "Active",
            },
            {
                "id": 2,
                "name": "Document Upload & OCR",
                "assigned_role": "System",
                "avg_duration_minutes": 2,
                "status": "Active",
            },
            {
                "id": 3,
                "name": "Bureau Check",
                "assigned_role": "System",
                "avg_duration_minutes": 1,
                "status": "Active",
            },
            {
                "id": 4,
                "name": "Credit Analysis",
                "assigned_role": "Credit Analyst",
                "avg_duration_minutes": 60,
                "status": "Active",
            },
            {
                "id": 5,
                "name": "Underwriting",
                "assigned_role": "Underwriter",
                "avg_duration_minutes": 120,
                "status": "Active",
            },
            {
                "id": 6,
                "name": "Compliance Review",
                "assigned_role": "Compliance Officer",
                "avg_duration_minutes": 30,
                "status": "Active",
            },
        ]

        total_minutes = sum(stage["avg_duration_minutes"] for stage in stages if stage["status"] == "Active")
        avg_completion_hours = round(total_minutes / 60, 1)

        return {
            "metrics": [
                {"key": "active_stages", "label": "Active Stages", "value": len([s for s in stages if s["status"] == "Active"])},
                {"key": "avg_completion_time", "label": "Avg. Completion Time", "value": f"{avg_completion_hours}h"},
                {"key": "sla_compliance", "label": "SLA Compliance", "value": "94.5%"},
                {"key": "active_workflows", "label": "Active Workflows", "value": 1},
            ],
            "workflow_name": "Standard Loan Workflow",
            "stages": stages,
            "conditions": [
                {
                    "id": "cond-001",
                    "condition": "If Credit Score < 650",
                    "outcome": "Route to Senior Underwriter for manual review",
                },
                {
                    "id": "cond-002",
                    "condition": "If DTI > 45%",
                    "outcome": "Require policy override approval",
                },
                {
                    "id": "cond-003",
                    "condition": "If Loan Amount > INR 20L",
                    "outcome": "Add additional compliance checks",
                },
            ],
        }

    def rule_engine_data(self) -> dict[str, Any]:
        rules = [
            {
                "id": 1,
                "name": "DTI Threshold Check",
                "category": "Financial",
                "condition": "DTI_RATIO > 45",
                "action": "Reject",
                "severity": "High",
                "status": "Active",
                "last_modified": "2026-04-01",
            },
            {
                "id": 2,
                "name": "Credit Score Minimum",
                "category": "Credit",
                "condition": "CIBIL_SCORE < 650",
                "action": "Flag for Review",
                "severity": "High",
                "status": "Active",
                "last_modified": "2026-03-28",
            },
            {
                "id": 3,
                "name": "Multiple Inquiries Check",
                "category": "Credit",
                "condition": "CREDIT_INQUIRIES_6M > 3",
                "action": "Flag for Review",
                "severity": "Medium",
                "status": "Active",
                "last_modified": "2026-04-05",
            },
            {
                "id": 4,
                "name": "LTV Limit",
                "category": "Financial",
                "condition": "LTV_RATIO > 80",
                "action": "Reject",
                "severity": "High",
                "status": "Active",
                "last_modified": "2026-03-15",
            },
            {
                "id": 5,
                "name": "Employment Stability",
                "category": "Income",
                "condition": "EMPLOYMENT_MONTHS < 12",
                "action": "Flag for Review",
                "severity": "Low",
                "status": "Inactive",
                "last_modified": "2026-02-20",
            },
        ]

        return {
            "metrics": [
                {"key": "total_rules", "label": "Total Rules", "value": len(rules)},
                {"key": "active_rules", "label": "Active Rules", "value": len([r for r in rules if r["status"] == "Active"])},
                {"key": "high_severity", "label": "High Severity", "value": len([r for r in rules if r["severity"] == "High"])},
                {"key": "triggered_today", "label": "Rules Triggered Today", "value": 12},
            ],
            "rules": rules,
        }


workflow_service = WorkflowService()